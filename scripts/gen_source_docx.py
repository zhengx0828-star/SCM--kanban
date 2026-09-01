import os, docx
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = r"C:\Users\ASUSONE\Desktop\scm-kanban"
OUT  = r"C:\Users\ASUSONE\Desktop\SCM供应链看板-完整源码.docx"

# ---------- 文件清单（有序） ----------
files = []
def add(p): files.append(p)

add("README.md")

add("backend/requirements.txt")
add("backend/app/__init__.py")
add("backend/app/database.py")
add("backend/app/models.py")
add("backend/app/schemas.py")
add("backend/app/main.py")
add("backend/app/seed.py")
for r in ["products","suppliers","materials","supply_relations","projects","rules"]:
    add(f"backend/app/routers/{r}.py")

for f in ["package.json","vite.config.ts","tsconfig.json","components.json","index.html"]:
    add(f"frontend/{f}")

for f in ["main.tsx","App.tsx","index.css"]:
    add(f"frontend/src/{f}")

for f in ["product","supplier","material","supplyRelation","project","rule"]:
    add(f"frontend/src/types/{f}.ts")
add("frontend/src/types/xlsx.d.ts")

for f in ["api.ts","utils.ts","theme.ts","cities.ts"]:
    add(f"frontend/src/lib/{f}")

for f in ["use-products.ts","use-suppliers.ts","use-materials.ts","use-supply-relations.ts","use-projects.ts","use-rules.ts"]:
    add(f"frontend/src/hooks/{f}")

for f in ["button.tsx","card.tsx","dialog.tsx","alert-dialog.tsx","input.tsx","label.tsx","select.tsx","skeleton.tsx","badge.tsx","textarea.tsx","checkbox.tsx","popover.tsx","command.tsx","table.tsx","pagination.tsx","sonner.tsx","empty-state.tsx","multi-select-filter.tsx"]:
    add(f"frontend/src/components/ui/{f}")

for f in ["SiteHeader.tsx","SiteSidebar.tsx"]:
    add(f"frontend/src/components/site/{f}")

for f in ["MaterialDrawerFilter.tsx","MaterialSupplySlicer.tsx","ProjectDemandSlicer.tsx","SupplyStudioTable.tsx","MaterialDetailDialog.tsx","ProjectDialogs.tsx","ExcelImportDialog.tsx"]:
    add(f"frontend/src/components/projects/{f}")

for f in ["SupplierDetailDialog.tsx","SupplierFormDialog.tsx","SupplierMasterFormDialog.tsx","SupplyRelationFormDialog.tsx"]:
    add(f"frontend/src/components/suppliers/{f}")

for f in ["DemandLineChart.tsx","MaterialSupplyCharts.tsx","ChinaMap.tsx"]:
    add(f"frontend/src/components/charts/{f}")

for f in ["ProductTable.tsx","ProductFormDialog.tsx","DeleteProductDialog.tsx","Pagination.tsx"]:
    add(f"frontend/src/components/{f}")

for f in ["DashboardPage.tsx","SupplyDemandPage.tsx","ProjectSupplyDetailPage.tsx","SupplierMasterPage.tsx","RulesPage.tsx","ProductsPage.tsx"]:
    add(f"frontend/src/pages/{f}")

add("start-all.bat")
add("start-all.sh")
for f in ["start-backend.bat","start-backend.sh","start-frontend.bat","start-frontend.sh"]:
    add(f"scripts/{f}")

# ---------- 文档生成 ----------
doc = Document()

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

for sec in doc.sections:
    sec.left_margin = Cm(1.8)
    sec.right_margin = Cm(1.8)

t = doc.add_heading("SCM 供应链看板 · 完整源码", level=0)
for run in t.runs:
    run.font.size = Pt(26)

p = doc.add_paragraph()
r = p.add_run("前后端分离 Web 应用：FastAPI + SQLAlchemy 2.0 + Pydantic v2（后端） × React 18 + TypeScript + Vite 5 + Tailwind CSS v4 + ECharts 6（前端）\n")
r.font.size = Pt(12)
p2 = doc.add_paragraph()
r2 = p2.add_run("本文档包含全部 80+ 个源码文件，供离线传输 / 内网环境恢复使用。恢复方式：按目录结构逐文件重建，或先在本机（有外网）用 opencode 复现后再用本文档逐文件核对覆盖。")
r2.font.size = Pt(11)
r2.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

doc.add_heading("项目目录结构", level=1)
tree = """scm-kanban/
├── start-all.bat / start-all.sh      # 一键启动（自动建 venv、装依赖、拉起前后端）
├── backend/                          # FastAPI 后端
│   ├── requirements.txt
│   ├── products.db                   # SQLite（启动自动建表，本文档不含数据）
│   ├── app/
│   │   ├── main.py / database.py / models.py / schemas.py / seed.py
│   │   └── routers/  (products / suppliers / materials / supply_relations / projects / rules)
│   └── scripts/  (smoke_test 等)
├── frontend/                         # React 前端
│   ├── package.json / vite.config.ts / tsconfig.json / index.html / components.json
│   ├── public/maps/china.json        # 中国地图边界数据（572KB，见文末说明，需单独拷贝）
│   └── src/
│       ├── main.tsx / App.tsx / index.css
│       ├── pages/     (Dashboard / SupplyDemand / ProjectSupplyDetail / SupplierMaster / Rules / Products)
│       ├── components/ (site / projects / suppliers / charts / ui / 产品演示组件)
│       ├── hooks/     (TanStack Query 封装 ×6)
│       ├── lib/       (api.ts / utils.ts / theme.ts / cities.ts)
│       └── types/     (7 个类型定义)
├── docs/                             # 截图 + 复现提示词（见文末说明）
└── scripts/                          # 前后端单独启动脚本
"""
for line in tree.rstrip("\n").split("\n"):
    tp = doc.add_paragraph()
    tr = tp.add_run(line)
    tr.font.name = "Consolas"
    tr.font.size = Pt(8.5)

doc.add_heading("使用说明（内网恢复三步）", level=1)
notes = [
    "1. 按本文档「目录结构」在目标电脑上逐目录重建文件，文件名与路径保持完全一致；",
    "2. 复制每个文件「源码」节中的完整内容覆盖对应文件；",
    "3. 双击 start-all.bat 一键启动（会自动创建 venv、安装依赖、拉起 5173/8000 服务）。",
    "",
    "注意：frontend/public/maps/china.json（572KB 地图数据）与 docs/ 截图未包含在本文档中。",
    "   - china.json 可让有外网的同事单独传一份，或从 opencode 复现的版本中取用（两者一致即可）；",
    "   - 没有它 Dashboard 地图不显示，其余功能不受影响。",
]
for n in notes:
    np_ = doc.add_paragraph()
    nr = np_.add_run(n)
    nr.font.size = Pt(10)

skipped = []
written = 0

for rel in files:
    full = os.path.join(ROOT, rel)
    if not os.path.isfile(full):
        skipped.append(rel)
        continue
    try:
        with open(full, "r", encoding="utf-8") as fh:
            content = fh.read()
    except Exception as e:
        skipped.append(f"{rel} ({e})")
        continue

    doc.add_heading(rel, level=1)
    meta = doc.add_paragraph()
    mr = meta.add_run(f"— 共 {content.count(chr(10))+1} 行 —")
    mr.font.size = Pt(9)
    mr.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    for line in content.rstrip("\n").split("\n"):
        cp = doc.add_paragraph()
        cp.paragraph_format.space_after = Pt(0)
        cp.paragraph_format.space_before = Pt(0)
        cp.paragraph_format.line_spacing = 1.0
        cr = cp.add_run(line if line else " ")
        cr.font.name = "Consolas"
        cr.font.size = Pt(8)
        cr.font.color.rgb = RGBColor(0x1a, 0x1a, 0x2e)
        pPr = cp._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "F5F5F5")
        pPr.append(shd)
    written += 1

doc.add_heading("未包含文件（需单独处理）", level=1)
tail = [
    "frontend/public/maps/china.json — 中国地图边界数据（572KB），Dashboard 地图必需，请单独拷贝；",
    "frontend/public/test-import.xlsx / xlsx.full.min.js — Excel 导入功能的测试文件与第三方库（可选）；",
    "backend/products.db — SQLite 数据文件，首次启动自动建空白表；",
    "node_modules / .venv / dist / __pycache__ — 依赖与构建产物，由 start-all.bat 自动生成。",
]
for t_ in tail:
    tp_ = doc.add_paragraph()
    tr_ = tp_.add_run(t_)
    tr_.font.size = Pt(10)

doc.save(OUT)
print(f"OK 已生成：{OUT}")
print(f"包含 {written} 个文件；跳过 {len(skipped)} 个：{skipped if skipped else '无'}")
print(f"大小：{os.path.getsize(OUT)/1024/1024:.2f} MB")
